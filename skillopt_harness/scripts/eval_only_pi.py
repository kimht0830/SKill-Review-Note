"""eval_only.py wrapper that injects privileged information (PI) into the target's user turn.

The gold answer of each item is prepended to the user prompt (the system prompt
is left untouched, so it stays identical across items). The repo files are not
edited: the per-env prompt builders are monkeypatched in this process and then
scripts/eval_only.py runs unchanged, so scoring/output layout are the same as the
no-skill baseline.

Templates come from the JSON file named by $SKILLOPT_PI_CONFIG (default: <harness>/pi_config.json):
    {"templates": {"<env>": "... {answer} ..."}, "spreadsheet_max_cells": 100}
"""
from __future__ import annotations

import atexit
import functools
import json
import os
import random
import runpy
import sys
import threading
import zlib

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.dirname(_SCRIPT_DIR)
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)


def _argv_value(flag: str) -> str:
    for i, arg in enumerate(sys.argv):
        if arg == flag and i + 1 < len(sys.argv):
            return sys.argv[i + 1]
        if arg.startswith(flag + "="):
            return arg.split("=", 1)[1]
    return ""


_CFG_PATH = os.environ.get("SKILLOPT_PI_CONFIG") or os.path.join(_PROJECT_ROOT, "pi_config.json")
if not _CFG_PATH or not os.path.isfile(_CFG_PATH):
    raise SystemExit("SKILLOPT_PI_CONFIG must point to the PI template JSON (refusing to run without PI).")
with open(_CFG_PATH, encoding="utf-8") as _f:
    _PI = json.load(_f)
_TEMPLATES: dict[str, str] = _PI["templates"]
_MAX_CELLS = int(_PI.get("spreadsheet_max_cells", 100))

# configs/<env>/default.yaml -> <env>
_ENV = os.path.basename(os.path.dirname(os.path.abspath(_argv_value("--config"))))
_OUT_ROOT = os.path.abspath(_argv_value("--out_root")) if _argv_value("--out_root") else ""

_lock = threading.Lock()
_seen: set[str] = set()
_stats = {"env": _ENV, "injected": 0, "skipped_no_reference": 0}


def _inject(key: str, reference: str, user: str) -> str:
    """Return the user prompt with the PI block prepended (unchanged if there is no reference)."""
    key = str(key)
    reference = str(reference or "").strip()
    if not reference:
        with _lock:
            if key not in _seen:
                _seen.add(key)
                _stats["skipped_no_reference"] += 1
        return user
    out = _TEMPLATES[_ENV].format(answer=reference).rstrip() + "\n\n" + user
    with _lock:
        if key not in _seen:  # some envs build the same prompt twice per item
            _seen.add(key)
            _stats["injected"] += 1
            if _OUT_ROOT:
                os.makedirs(_OUT_ROOT, exist_ok=True)
                with open(os.path.join(_OUT_ROOT, "pi_injections.jsonl"), "a", encoding="utf-8") as f:
                    f.write(json.dumps({"id": key, "reference": reference}, ensure_ascii=False) + "\n")
                if _stats["injected"] == 1:
                    with open(os.path.join(_OUT_ROOT, "pi_example_user_prompt.txt"), "w", encoding="utf-8") as f:
                        f.write(out)
    return out


def _require(module, *names: str) -> None:
    missing = [n for n in names if not callable(getattr(module, n, None))]
    if missing:
        raise SystemExit(
            f"[PI] {module.__name__} has no {missing}: the SkillOpt repo layout changed, "
            "update the patches in scripts/eval_only_pi.py."
        )


# ── per-env patches ─────────────────────────────────────────────────────────

def _patch_searchqa() -> None:
    from skillopt.envs.searchqa import rollout as m

    _require(m, "_build_user", "process_one")
    orig_build, orig_process = m._build_user, m.process_one
    current = threading.local()  # _build_user only receives question/context, not the item

    @functools.wraps(orig_process)
    def process_one(item, *args, **kwargs):
        current.item = item
        try:
            return orig_process(item, *args, **kwargs)
        finally:
            current.item = None

    @functools.wraps(orig_build)
    def _build_user(*args, **kwargs):
        user = orig_build(*args, **kwargs)
        item = getattr(current, "item", None) or {}
        answers = [str(a) for a in (item.get("answers") or []) if str(a).strip()]
        return _inject(item.get("id", ""), answers[0] if answers else "", user)

    m.process_one, m._build_user = process_one, _build_user


def _patch_livemathematicianbench() -> None:
    from skillopt.envs.livemathematicianbench import rollout as m

    _require(m, "_build_user")
    orig_build = m._build_user

    @functools.wraps(orig_build)
    def _build_user(item, *args, **kwargs):
        user = orig_build(item, *args, **kwargs)
        correct = item.get("correct_choice") or {}  # already shuffled by the dataloader
        label, text = str(correct.get("label", "")).strip(), str(correct.get("text", "")).strip()
        return _inject(item.get("id", ""), f"{label}. {text}" if label else "", user)

    m._build_user = _build_user


def _patch_docvqa() -> None:
    from skillopt.envs.docvqa import rollout as m
    from skillopt.envs.docvqa.evaluator import _extract_answer_strings

    _require(m, "_build_messages")
    orig_build = m._build_messages

    @functools.wraps(orig_build)
    def _build_messages(item, *args, **kwargs):
        messages, system, user_text = orig_build(item, *args, **kwargs)
        answers = [a for a in _extract_answer_strings(item.get("answers", [])) if str(a).strip()]
        new_text = _inject(item.get("id", ""), answers[0] if answers else "", user_text)
        for part in messages[1]["content"]:
            if part.get("type") == "text":
                part["text"] = new_text
                break
        else:
            raise RuntimeError("[PI] docvqa user message has no text part")
        return messages, system, new_text

    m._build_messages = _build_messages


def _patch_officeqa() -> None:
    from skillopt.envs.officeqa import rollout as m

    _require(m, "_build_user")
    orig_build = m._build_user

    @functools.wraps(orig_build)
    def _build_user(item, *args, **kwargs):
        user = orig_build(item, *args, **kwargs)
        reference = item.get("ground_truth") or next(iter(item.get("answers") or []), "")
        return _inject(item.get("id", ""), reference, user)

    m._build_user = _build_user


def _spreadsheet_gold_path(input_xlsx: str) -> str:
    folder, name = os.path.split(input_xlsx)
    for src, dst in (("_input.xlsx", "_answer.xlsx"), ("_init.xlsx", "_golden.xlsx")):
        if name.endswith(src):
            return os.path.join(folder, name[: -len(src)] + dst)
    if name == "initial.xlsx":
        return os.path.join(folder, "golden.xlsx")
    return ""


def _fmt_cell(value) -> str:
    if value is None:
        return "(empty)"
    return repr(value) if isinstance(value, str) else str(value)


@functools.lru_cache(maxsize=None)
def _spreadsheet_expected_cells(gold_path: str, answer_position: str) -> str:
    """Gold values of the answer cells. Ranges with more than _MAX_CELLS cells are shown as a
    random sample of _MAX_CELLS cells (seeded per item, so reruns show the same cells)."""
    import openpyxl

    from skillopt.envs.spreadsheetbench.evaluator import _generate_cell_names

    try:
        wb = openpyxl.load_workbook(filename=gold_path, data_only=True)
    except Exception:  # noqa: BLE001
        return ""
    try:
        cells: list[tuple[str, str]] = []
        for scr in (answer_position or "").split(","):
            scr = scr.strip()
            if not scr:
                continue
            if "!" in scr:
                sheet, cell_range = scr.split("!", 1)
                sheet = sheet.strip().strip("'\"")
            else:
                sheet, cell_range = wb.sheetnames[0], scr
            if sheet not in wb.sheetnames:
                continue
            cells += [(sheet, name) for name in _generate_cell_names(cell_range.strip().strip("'\""))]
        total = len(cells)
        if not total:
            return ""
        lines: list[str] = []
        if total > _MAX_CELLS:
            item_key = "/".join(gold_path.replace("\\", "/").split("/")[-2:])  # <task dir>/<gold file>
            rng = random.Random(zlib.crc32(f"{item_key}|{answer_position}".encode("utf-8")))
            cells = [cells[i] for i in sorted(rng.sample(range(total), _MAX_CELLS))]
            lines.append(
                f"[TRUNCATED] The answer range contains {total} cells. Only {_MAX_CELLS} randomly selected "
                f"cells are listed below; the other {total - _MAX_CELLS} answer cells also have required "
                "values that are not shown."
            )
        lines += [f"{sheet}!{name} = {_fmt_cell(wb[sheet][name].value)}" for sheet, name in cells]
        return "\n".join(lines)
    except Exception:  # noqa: BLE001 - unparsable range: the evaluator fails on it too
        return ""
    finally:
        wb.close()


def _patch_spreadsheetbench() -> None:
    from skillopt.envs.spreadsheetbench import codegen_agent as m

    _require(m, "_build_user")
    orig_build = m._build_user

    @functools.wraps(orig_build)
    def _build_user(instruction, input_xlsx, instruction_type="", answer_position="", **kwargs):
        user = orig_build(instruction, input_xlsx, instruction_type, answer_position, **kwargs)
        gold = _spreadsheet_gold_path(input_xlsx)
        cells = _spreadsheet_expected_cells(gold, answer_position) if gold and os.path.exists(gold) else ""
        return _inject(input_xlsx, cells, user)

    m._build_user = _build_user


_PATCHES = {
    "searchqa": _patch_searchqa,
    "livemathematicianbench": _patch_livemathematicianbench,
    "docvqa": _patch_docvqa,
    "officeqa": _patch_officeqa,
    "spreadsheetbench": _patch_spreadsheetbench,
}


def _report() -> None:
    print(f"  [PI] env={_ENV} injected={_stats['injected']} "
          f"skipped_no_reference={_stats['skipped_no_reference']}", flush=True)
    if not _stats["injected"]:
        print("  [PI] WARNING: no prompt received privileged information in this run "
              "(expected only when every item was resumed from an existing results.jsonl).", flush=True)
    if _OUT_ROOT and os.path.isdir(_OUT_ROOT):
        with open(os.path.join(_OUT_ROOT, "pi_summary.json"), "w", encoding="utf-8") as f:
            json.dump(_stats, f, indent=2)


if __name__ == "__main__":
    if _ENV not in _PATCHES or _ENV not in _TEMPLATES:
        raise SystemExit(f"[PI] no PI patch/template for env {_ENV!r} (supported: {sorted(_PATCHES)})")
    _PATCHES[_ENV]()
    atexit.register(_report)
    print(f"  [PI] privileged-information injection enabled for env={_ENV}", flush=True)
    runpy.run_path(os.path.join(_SCRIPT_DIR, "eval_only.py"), run_name="__main__")